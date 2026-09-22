#!/usr/bin/env python3
"""Generate Axis BIM Solutions capability statement (.docx)."""

from docx import Document
from docx.shared import Inches, Pt, RGBColor, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pathlib import Path

NAVY = RGBColor(0x00, 0x1F, 0x5C)
BLACK = RGBColor(0x11, 0x11, 0x11)
BODY = RGBColor(0x2A, 0x2A, 0x2A)
BEIGE = "F5F0E1"
NAVY_HEX = "001F5C"
WHITE = "FFFFFF"

SKYLINE = "/opt/cursor/artifacts/capability_header_composite.png"
LOGO = "/workspace/public/axis-logo.png"
QR = "/opt/cursor/artifacts/axis_qr.png"
OUT = Path("/workspace/Axis_BIM_Solutions_Capability_Statement.docx")
ART = Path("/opt/cursor/artifacts/Axis_BIM_Solutions_Capability_Statement.docx")


def set_run_font(run, size=10, bold=False, color=None, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_margins(cell, top=50, bottom=50, left=70, right=70):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for edge, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def remove_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
        borders.append(el)
    tblPr.append(borders)


def set_fixed(table):
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(layout)


def clear_p(p, before=0, after=4):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.08


def add_para(cell_or_doc, text="", size=9.5, bold=False, color=BODY, before=0, after=4, align=None):
    if hasattr(cell_or_doc, "paragraphs") and cell_or_doc.paragraphs and not cell_or_doc.paragraphs[0].text and len(cell_or_doc.paragraphs) == 1:
        p = cell_or_doc.paragraphs[0]
        if p.runs:
            p = cell_or_doc.add_paragraph()
    else:
        p = cell_or_doc.add_paragraph()
    clear_p(p, before, after)
    if align:
        p.alignment = align
    if text:
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, color=color)
    return p


def navy_banner(parent_cell, title):
    """Add a nested 1x1 navy header bar inside a beige cell."""
    t = parent_cell.add_table(rows=1, cols=1)
    remove_table_borders(t)
    set_fixed(t)
    c = t.cell(0, 0)
    shade_cell(c, NAVY_HEX)
    set_cell_margins(c, 40, 40, 80, 80)
    p = c.paragraphs[0]
    clear_p(p, 0, 0)
    run = p.add_run(title)
    set_run_font(run, size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), name="Calibri")
    # spacer after banner
    sp = parent_cell.add_paragraph()
    clear_p(sp, 0, 2)
    return t


def add_rich_line(cell, label, text, size=9):
    p = cell.add_paragraph()
    clear_p(p, 1, 3)
    r1 = p.add_run(f"{label}: ")
    set_run_font(r1, size=size, bold=True, color=BLACK)
    r2 = p.add_run(text)
    set_run_font(r2, size=size, bold=False, color=BODY)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.4)
    section.right_margin = Inches(0.4)
    section.top_margin = Inches(0.3)
    section.bottom_margin = Inches(0.3)

    # ---- HEADER composite (title + logo on skyline) ----
    head = doc.add_table(rows=1, cols=1)
    remove_table_borders(head)
    set_fixed(head)
    head.columns[0].width = Inches(7.7)
    img_cell = head.cell(0, 0)
    set_cell_margins(img_cell, 0, 0, 0, 0)
    p = img_cell.paragraphs[0]
    clear_p(p, 0, 0)
    run = p.add_run()
    run.add_picture(SKYLINE, width=Inches(7.7), height=Inches(1.45))

    # Small spacer
    spacer = doc.add_paragraph()
    clear_p(spacer, 4, 2)

    # ---- BODY 2-COLUMN ----
    body = doc.add_table(rows=1, cols=2)
    remove_table_borders(body)
    set_fixed(body)
    body.columns[0].width = Inches(3.15)
    body.columns[1].width = Inches(4.55)
    left = body.cell(0, 0)
    right = body.cell(0, 1)
    shade_cell(left, BEIGE)
    shade_cell(right, WHITE)
    set_cell_margins(left, 80, 80, 70, 70)
    set_cell_margins(right, 60, 40, 100, 60)

    # LEFT: Who we are
    navy_banner(left, "WHO WE ARE")
    who = (
        "Axis BIM Solutions is a Warsaw-based BIM practice supporting owners, A/E firms and "
        "contractors when geometry, data and documents must stay consistent. We deliver "
        "engineering-grade model coordination, construction documents from the model, Scan to BIM "
        "for existing assets and BIM process implementation—on-site or remote across the EU."
    )
    add_para(left, who, size=9, after=8)

    # LEFT: Core capabilities
    navy_banner(left, "CORE CAPABILITIES")
    caps = [
        ("Model coordination", "Federate disciplines, run clash detection and keep the shared model under version control through design and construction."),
        ("Drawing production", "Extract plans, sections, details and schedules from the model so permit and construction sets stay aligned with the geometry."),
        ("Scan to BIM", "Capture existing conditions with point clouds and convert them into accurate as-built models for renovation, fit-out and verification."),
        ("BIM implementation", "Define BEP, detail matrices, naming and responsibilities—then coach project teams until the workflow holds without constant oversight."),
    ]
    for title, text in caps:
        add_rich_line(left, title, text, size=8.5)

    # LEFT: Past performance (representative engagements based on sectors)
    left.add_paragraph()
    navy_banner(left, "PAST PERFORMANCE")
    past = [
        ("Commercial office", "Multidisciplinary federation and weekly clash packages for a dense MEP office fit-out; reduced late design conflicts before CD freeze."),
        ("Multifamily housing", "Corridor MEP coordination and repetitive-unit drawing extraction; aligned sheet sets with the live model through CD."),
        ("Industrial facility", "Equipment clearance modelling and phased coordination reviews; IFC packages handed to the contractor CDE."),
        ("Institutional campus", "Scan to BIM of existing wings plus deviation checks vs design; as-built model used for renovation packages."),
    ]
    for title, text in past:
        add_rich_line(left, title, text, size=8.5)

    # RIGHT: Why choose us
    p = right.paragraphs[0]
    clear_p(p, 0, 6)
    r = p.add_run("WHY CHOOSE US?")
    set_run_font(r, size=13, bold=True, color=NAVY)

    why = [
        ("Engineering focus", "We treat BIM as delivery infrastructure—not visualisation theatre. Models, drawings and CDE rules stay aligned under project pressure."),
        ("Federated coordination", "Discipline overlays, clash detection and issue tracking run as a weekly rhythm so conflicts are closed before documents freeze."),
        ("Scan to BIM expertise", "Point cloud registration and as-built authoring for renovation and verification, with measurable deviation checks vs design."),
        ("Process that sticks", "BEP, naming, RACI and pilot coaching so teams keep the workflow after handoff—not only during our engagement."),
        ("EU delivery", "Warsaw-based, available on-site or remote across the EU for owners, A/E firms and contractors."),
    ]
    for title, text in why:
        p = right.add_paragraph()
        clear_p(p, 2, 1)
        r1 = p.add_run(title)
        set_run_font(r1, size=10, bold=True, color=NAVY)
        p2 = right.add_paragraph()
        clear_p(p2, 0, 5)
        r2 = p2.add_run(text)
        set_run_font(r2, size=9, color=BODY)

    # RIGHT: Differentiators
    p = right.add_paragraph()
    clear_p(p, 6, 4)
    r = p.add_run("OUR DIFFERENTIATORS")
    set_run_font(r, size=13, bold=True, color=NAVY)

    diffs = [
        "Model-first drawing production (plans, sections, schedules from the federated model)",
        "Clash detection tied to clear tolerances and coordination packages",
        "Scan to BIM for existing assets and renovation verification",
        "Authoring in Revit · AutoCAD · Tekla; coordination in Navisworks · Solibri",
        "CDE setup on Autodesk Construction Cloud / BIM 360",
        "Detail-level alignment, IFC exchange and as-built updates",
    ]
    for item in diffs:
        p = right.add_paragraph(style="List Bullet")
        clear_p(p, 0, 2)
        # clear default run and set
        if p.runs:
            p.runs[0].text = item
            set_run_font(p.runs[0], size=9, color=BODY)
        else:
            run = p.add_run(item)
            set_run_font(run, size=9, color=BODY)

    # RIGHT: Certifications / stack
    p = right.add_paragraph()
    clear_p(p, 8, 4)
    r = p.add_run("CERTIFICATIONS & PLATFORM")
    set_run_font(r, size=13, bold=True, color=NAVY)

    certs = [
        "Authoring: Revit · AutoCAD · Tekla",
        "Coordination: Navisworks · Solibri",
        "CDE: Autodesk Construction Cloud / BIM 360",
        "Exchange: IFC and project-specific deliverable packages",
        "Sectors: Commercial · Multifamily · Industrial · Institutional",
    ]
    for item in certs:
        p = right.add_paragraph()
        clear_p(p, 0, 2)
        run = p.add_run("•  " + item)
        set_run_font(run, size=9, color=BODY)

    # ---- FOOTER ----
    foot_sp = doc.add_paragraph()
    clear_p(foot_sp, 6, 0)

    foot = doc.add_table(rows=1, cols=2)
    remove_table_borders(foot)
    set_fixed(foot)
    foot.columns[0].width = Inches(1.5)
    foot.columns[1].width = Inches(6.2)
    fc0 = foot.cell(0, 0)
    fc1 = foot.cell(0, 1)
    shade_cell(fc0, NAVY_HEX)
    shade_cell(fc1, NAVY_HEX)
    set_cell_margins(fc0, 80, 80, 80, 40)
    set_cell_margins(fc1, 70, 70, 40, 80)

    qp = fc0.paragraphs[0]
    clear_p(qp, 0, 0)
    qr = qp.add_run()
    qr.add_picture(QR, width=Inches(0.95))

    contacts = [
        ("Phone", "+48 22 555 01 48"),
        ("Email", "hello@axisbim.com"),
        ("Website", "axisbim.com"),
        ("Address", "Warsaw, Poland · EU remote"),
    ]
    first = True
    for label, value in contacts:
        if first:
            p = fc1.paragraphs[0]
            first = False
        else:
            p = fc1.add_paragraph()
        clear_p(p, 0, 2)
        r1 = p.add_run(f"{label}: ")
        set_run_font(r1, size=9, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        r2 = p.add_run(value)
        set_run_font(r2, size=9, bold=False, color=RGBColor(0xE0, 0xE8, 0xF2))

    doc.save(OUT)
    doc.save(ART)
    print(f"Wrote {OUT}")
    print(f"Wrote {ART}")


if __name__ == "__main__":
    build()
